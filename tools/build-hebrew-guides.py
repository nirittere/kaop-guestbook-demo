from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables"
OUTPUT.mkdir(parents=True, exist_ok=True)

FONT = "Arial"
NAVY = "17365D"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F9FC"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = RGBColor(72, 72, 72)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_run_font(run, size=None, bold=None, color=None, rtl=None):
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:cs"), FONT)
    if rtl is not None:
        rtl_node = r_pr.find(qn("w:rtl"))
        if rtl_node is None:
            rtl_node = OxmlElement("w:rtl")
            r_pr.append(rtl_node)
        rtl_node.set(qn("w:val"), "1" if rtl else "0")
        lang = r_pr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            r_pr.append(lang)
        if rtl:
            lang.set(qn("w:bidi"), "he-IL")
        else:
            lang.set(qn("w:val"), "en-US")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_ltr(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "0")


def set_style_rtl(style, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    style.paragraph_format.alignment = alignment
    p_pr = style._element.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_section_rtl(section):
    sect_pr = section._sectPr
    bidi = sect_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        sect_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bidi_visual = tbl_pr.find(qn("w:bidiVisual"))
    if bidi_visual is None:
        bidi_visual = OxmlElement("w:bidiVisual")
        tbl_pr.insert(0, bidi_visual)
    bidi_visual.set(qn("w:val"), "1")


def set_rtl_list_indent(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:right"), "420")
    ind.set(qn("w:hanging"), "260")


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, 9, color=TEXT_GRAY, rtl=False)


def configure_document(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    set_section_rtl(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    set_style_rtl(normal)

    title_style = styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title_style._element.rPr.rFonts.set(qn("w:cs"), FONT)
    title_style.font.size = Pt(28)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    set_style_rtl(title_style)
    remove_paragraph_border(title_style)

    for style_name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        set_style_rtl(style)

    p = doc.add_paragraph(style="Title")
    set_rtl(p)
    remove_paragraph_border(p)
    set_run_font(p.add_run(title), rtl=True)
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(20)

    p = doc.add_paragraph()
    set_rtl(p)
    run = p.add_run(subtitle)
    set_run_font(run, 15, color=TEXT_GRAY, rtl=True)
    p.paragraph_format.space_after = Pt(32)

    p = doc.add_paragraph()
    set_rtl(p)
    run = p.add_run("KAOP by Nirit Terehovsky")
    set_run_font(run, 13, bold=True, rtl=False)

    p = doc.add_paragraph()
    set_rtl(p)
    run = p.add_run("גרסה מעודכנת על בסיס הדמו המאומת מיום 7 באוקטובר 2026")
    set_run_font(run, 10.5, color=TEXT_GRAY, rtl=True)

    doc.add_page_break()
    add_footer(doc)
    return doc


def add_footer(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        add_page_number(p)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_rtl(p)
    set_run_font(p.add_run(text), rtl=True)
    set_keep_with_next(p)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    set_rtl(p)
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, bold=True, rtl=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, rtl=True)
    else:
        run = p.add_run(text)
        set_run_font(run, rtl=True)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        set_rtl(p)
        set_rtl_list_indent(p)
        run = p.add_run(f"• {item}")
        set_run_font(run, rtl=True)
        p.paragraph_format.space_after = Pt(4)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        set_rtl(p)
        set_rtl_list_indent(p)
        run = p.add_run(f"{index}. {item}")
        set_run_font(run, rtl=True)
        p.paragraph_format.space_after = Pt(5)


def add_code(doc, text):
    p = doc.add_paragraph()
    set_ltr(p)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    run = p.add_run(text)
    set_run_font(run, 9.5, rtl=False)
    run.font.name = "Courier New"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Courier New")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Courier New")
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_rtl(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    prevent_row_split(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(header)
        set_run_font(run, 10.5, bold=True, color=RGBColor(255, 255, 255), rtl=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            set_rtl(p)
            run = p.add_run(str(value))
            set_run_font(run, 10, rtl=True)
    set_table_borders(table)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph()
    return table


def add_link_lines(doc, links):
    for label, url in links:
        p = doc.add_paragraph()
        set_ltr(p)
        p.paragraph_format.space_after = Pt(5)
        lead = p.add_run(f"{label}: ")
        set_run_font(lead, 10.5, bold=True, rtl=False)
        run = p.add_run(url)
        set_run_font(run, 10.5, rtl=False)


def add_contents(doc, items):
    add_heading(doc, "תוכן המסמך", 1)
    add_numbered(doc, items)
    doc.add_page_break()


def save(doc, filename):
    path = OUTPUT / filename
    doc.save(path)
    return path


def build_talk_guide():
    doc = configure_document(
        "מדריך להעברת ההרצאה על KAOP",
        "תסריט דיבור מלא למצגת ולמעבר לדמו",
    )
    add_contents(doc, [
        "מטרת ההרצאה והמסר המרכזי",
        "חלוקת הזמן",
        "תסריט לפי שקפים",
        "מעבר לדמו וסיום",
        "שאלות צפויות ותשובות",
    ])

    add_heading(doc, "מטרת ההרצאה והמסר המרכזי")
    add_para(doc, "המטרה שלי היא להראות כיצד KAOP נותנת לצוות Platform Engineering שכבת שליטה משותפת לסוכני AI שפועלים בסביבות ייצור. ההרצאה אינה מציגה אוסף תכונות. היא מציגה בעיה תפעולית, את מודל השליטה של KAOP, ולאחר מכן ראיות חיות מתוך incident אמיתי וסקירת קוד אמיתית.")
    add_para(doc, "המסר המרכזי הוא שאפשר להרחיב שימוש בסוכנים בלי לוותר על הרשאות מוגבלות, עקיבות, אחריות ובקרה אנושית. בדמו הסוכן של Kubernetes קורא מידע בלבד בתוך namespace אחד. הסוכן של GitHub קורא קוד וכותב תגובת review בלבד.")

    add_heading(doc, "חלוקת הזמן")
    add_table(doc, ["חלק", "זמן מומלץ", "מטרה"], [
        ["פתיחה ושקפים 1 עד 3", "3 דקות", "הגדרת בעיית agent sprawl והפער בין agent מקומי להפעלה בייצור"],
        ["שקפים 4 עד 7", "4 דקות", "הסבר על control plane, workflow, governance ו-run evidence"],
        ["שקפים 8 עד 9", "2 דקות", "הכנת הקהל לדמו ולמודל האימוץ"],
        ["שקפים 10 עד 12", "2 עד 3 דקות", "שקפי appendix לפי שאלות או כהכנה טכנית קצרה"],
        ["דמו חי", "עד 12 דקות", "אירוע, alert, RCA, review והתאוששות"],
    ], [4.2, 3.0, 9.0])

    slides = [
        (1, "KAOP by Nirit Terehovsky", "כדאי לפתוח במשפט הבא: צוותי פיתוח כבר משתמשים בסוכנים. השאלה של צוות הפלטפורמה היא כיצד להרחיב את השימוש בלי לאבד שליטה על הרשאות, ראיות, עלות ואחריות. היום אראה כיצד KAOP מספקת שכבת הפעלה משותפת לסוכנים בסביבת ייצור.", "המעבר: כדי להבין מדוע נדרשת שכבה כזאת, נתחיל במה שקורה כאשר כל צוות מאמץ סוכן באופן עצמאי."),
        (2, "The agent sprawl problem", "כל צוות יכול לבחור runtime, מודל, credentials ותהליך review משלו. ההתקדמות המקומית מהירה, אבל צוות הפלטפורמה מתקשה לענות על שאלות בסיסיות: אילו סוכנים פועלים, למי יש גישה, מי אישר פעולה, ואילו ראיות נשמרו.", "המעבר: גם agent מצוין מבחינה פונקציונלית עדיין חסר את מעטפת הייצור."),
        (3, "Why isolated agents fail in production", "יכולת reasoning לבדה אינה מספיקה. סוכן בסביבת ייצור צריך הרשאות מוגדרות, trigger אמין, workflow חוזר, ראיות לכלי העבודה שלו, ומסלול approval לפעולות מסוכנות. אחרת כל צוות בונה מעטפת שונה עם פערים אחרים.", "המעבר: KAOP מרכזת את המעטפת הזאת ב-control plane אחד."),
        (4, "The KAOP control plane", "KAOP מפרידה בין הסוכן לבין התשתית שמנהלת אותו. Fleet מציג inventory ובעלות. Run מטפל ב-triggers וב-workflows. Context מחבר ידע ואינטגרציות. Govern אוכף RBAC, credentials, guardrails ו-approvals. כך אפשר לנהל סוכנים ממקורות שונים תחת מודל הפעלה אחד.", "המעבר: נראה כיצד control plane הופך alert לתהליך חקירה מובנה."),
        (5, "Workflow orchestration from trigger to outcome", "ה-workflow מתחיל ב-Grafana alert, אוסף ראיות מ-Kubernetes, מחזיר RCA מובנה, משאיר החלטה מסוכנת לאדם ושומר run record. ה-output contract דורש root cause, evidence, confidence, affected resources ו-suggested fix. המבנה מאפשר לאנשים ולשלבים נוספים להשתמש בתוצאה.", "המעבר: ה-workflow שימושי רק אם ההרשאות שלו נאכפות בפועל."),
        (6, "Governance boundaries", "בצד Kubernetes יצרתי ServiceAccount שמוגבל ל-namespace kaop-demo ולפעולות get, list ו-watch. הוא אינו יכול לקרוא Secrets, לבצע write או לפתוח pods exec. בצד GitHub ה-reviewer יכול לקרוא repository ו-PR ולכתוב תגובת review, אך אינו יכול push או merge.", "המעבר: מעבר להגבלת גישה, צוות הפלטפורמה צריך יכולת לבדוק מה קרה בכל run."),
        (7, "Fleet visibility and auditability", "Run history מרכז ownership, inputs, tool calls, policy decisions ותוצאה. במקום לשחזר אירוע מתוך terminals ולוגים מפוזרים, צוות הפלטפורמה מקבל record אחד שניתן לבדיקה. בדמו נראה run אמיתי עם 15 tool calls ו-high confidence finding.", "המעבר: עכשיו אפשר לעבור ממודל תאורטי לראיות החיות."),
        (8, "Live platform tour", "הדמו מחבר שלוש מערכות. Grafana מזהה restarts ושולחת webhook. KAOP מפעילה RCA מוגבל הרשאות. GitHub מציג PR פתוח עם review של AgentOps. התקלה נוצרת באמצעות make incident ומסירה endpoints אמיתיים משירות Redis.", "המעבר: לפני הדמו, חשוב להסביר כיצד מרחיבים autonomy בצורה מבוקרת."),
        (9, "A governed path to agent adoption", "הגישה המומלצת מתחילה ב-read only ובתוצאה מדידה. לאחר שנצברות ראיות אפשר לעבור להמלצות, להצעות write עם approval, ורק לבסוף לאוטומציה צרה. בדמו אנו נשארים בשלב Observe כדי לשמור על blast radius ברור.", "המעבר: בשלב זה עוברים לדמו. שקפים 10 עד 12 זמינים לדיון טכני או כ-fallback."),
        (10, "Demo architecture", "Grafana Cloud מקבלת metrics מה-cluster. alert שולח webhook ל-KAOP. סוכן RCA קורא את namespace kaop-demo בלבד. GitHub מחזיק את PR התיקון. אותו incident מחבר את החקירה ואת סקירת הקוד, ולכן הדמו נשאר סיפור אחד ולא feature tour מפוצל.", "המעבר: אם עולה שאלה על הרשאות, עוברים למטריצה בשקף הבא."),
        (11, "Permission matrix", "המילה Denied היא חלק מההוכחה. היא מציינת שהגישה ל-Secrets ולפעולות write חסומה במפורש באמצעות RBAC. אין מדובר בכשל חיבור. זהו גבול אבטחה שנבדק עם kubectl auth can-i. גם GitHub מוגבל ל-read ולכתיבת comment.", "המעבר: השקף האחרון מסכם את ההבדל בין מצב תקין, incident והתאוששות."),
        (12, "Incident evidence and recovery", "במצב תקין יש Redis endpoint ושני frontend Pods מוכנים. ב-incident ה-selector אינו תואם ל-Pods, אין endpoint, ה-Pod החדש אינו Ready ונצפו שבעה restarts. לאחר make recover ה-endpoint חוזר, ה-Pods מוכנים ו-Grafana חוזרת ל-Normal.", "המעבר: מכאן מסיימים במסר העסקי או חוזרים למסכי ה-run לצורך שאלות."),
    ]
    add_heading(doc, "תסריט לפי שקפים")
    for number, title, script, transition in slides:
        add_heading(doc, f"שקף {number} {title}", 2)
        add_para(doc, script)
        add_para(doc, transition, "המעבר:")

    add_heading(doc, "מעבר לדמו וסיום")
    add_para(doc, "משפט מעבר מומלץ: כעת אראה את אותו מודל בפעולה. נתחיל מההרשאות ומה-run history, ניצור תקלה אמיתית, נראה את Grafana מפעילה RCA, נבדוק את ה-review ב-GitHub ולבסוף נחזיר את המערכת למצב תקין.")
    add_para(doc, "משפט סיום מומלץ: הערך של KAOP הוא שכבת ההפעלה סביב הסוכנים. צוות הפלטפורמה מקבל מקום אחד להריץ workflows, לאכוף גבולות ולשמור ראיות. צוותי המוצר יכולים להמשיך לאמץ סוכנים בלי להפוך גישה והרשאות לתהליך בלתי נראה.")

    add_heading(doc, "שאלות צפויות ותשובות")
    qa = [
        ("למה לא לאפשר remediation אוטומטי", "המטרה בדמו היא להוכיח שליטה וחקירה לפני הרחבת autonomy. write אוטומטי יגדיל את הסיכון ויסיט את הדיון מה-governance."),
        ("האם KAOP קוראת Secrets", "לא. ה-Role אינו כולל את resource secrets. בדיקת kubectl auth can-i מחזירה no."),
        ("למה ה-liveness probe נפרד מ-readiness", "readiness יכולה לשקף תלות ב-Redis ולהוציא Pod מה-Service. liveness צריכה לבדוק אם התהליך עצמו חי. חיבור liveness לתלות חיצונית יוצר restart loop מיותר."),
        ("האם ה-alert באמת הפעיל את KAOP", "כן. Grafana שלחה webhook לאחר מעבר ל-Firing. ה-run המאומת הוא run_e15c91fc04701d47c908ec0f."),
        ("האם reviewer יכול לשנות קוד", "לא. האינטגרציה מספקת קריאת תוכן ו-PR וכתיבת review בלבד. ה-PR נשאר פתוח."),
        ("מה מייצג Denied בשקף 11", "Denied מציין פעולה שנחסמה בכוונה על ידי מנגנון ההרשאות. זו הוכחת governance, לא סטטוס תקלה."),
    ]
    add_table(doc, ["שאלה", "תשובה קצרה"], qa, [6.2, 10.0])

    return save(doc, "01-KAOP-presentation-talk-guide-he.docx")


def build_system_guide():
    doc = configure_document(
        "מדריך המערכת של הדמו KAOP",
        "הרכיבים האחריות של כל רכיב והחיבורים ביניהם",
    )
    add_contents(doc, [
        "תמונת המערכת",
        "האפליקציה והתקלה",
        "Kubernetes והרשאות",
        "Grafana Cloud וה-alert",
        "KAOP ו-workflow החקירה",
        "GitHub ו-Code Reviewer",
        "חומרי ההגשה",
        "זרימת המידע והאבטחה",
    ])

    add_heading(doc, "תמונת המערכת")
    add_para(doc, "המערכת בנויה סביב אפליקציית Guestbook ב-Go שתלויה ב-Redis. האפליקציה רצה ב-cluster מקומי מסוג k3d. Grafana Cloud אוספת metrics ומזהה restarts. KAOP מקבלת webhook, מפעילה Kubernetes RCA agent מוגבל הרשאות ושומרת את תוצאת החקירה. GitHub מחזיק את קוד הדמו ואת PR התיקון.")
    add_table(doc, ["רכיב", "תפקיד", "אחריות עיקרית"], [
        ["Go Guestbook", "Frontend", "מגיש ממשק HTTP, בודק חיבור ל-Redis ומספק livez ו-readyz"],
        ["Redis", "Data service", "שומר את נתוני ה-Guestbook ומספק dependency אמיתית"],
        ["k3d ו-Kubernetes", "Runtime", "מריץ Pods, Services, probes, rollouts ו-RBAC"],
        ["Kustomize", "Manifest composition", "מפריד בין מצב healthy לבין תרחיש incident"],
        ["Makefile", "Operator interface", "מספק פקודות יציבות ל-bootstrap, incident, verify ו-recover"],
        ["Grafana Alloy", "Telemetry collector", "מעביר metrics מה-cluster אל Grafana Cloud"],
        ["kube-state-metrics", "Kubernetes metrics", "חושף זמינות workloads ומוני restart"],
        ["Grafana Alerting", "Detection", "מעריך את תנאי ה-alert ושולח webhook ל-KAOP"],
        ["KAOP RCA agent", "Investigation", "אוסף ראיות מ-Kubernetes ומחזיר RCA מובנה"],
        ["GitHub reviewer", "Code review", "קורא את PR התיקון וכותב review comment"],
    ], [3.5, 4.0, 8.7])

    add_heading(doc, "האפליקציה והתקלה")
    add_heading(doc, "Go Guestbook", 2)
    add_para(doc, "הקובץ main.go מפעיל HTTP server ומתחבר ל-Redis דרך REDIS_ADDRESS. הנתיב livez בודק שה-process עצמו פעיל. הנתיב readyz בודק גם את התלות ב-Redis. ההפרדה חשובה משום שכשל בתלות צריך לעצור traffic אל ה-Pod, אך לא בהכרח להפעיל restart של התהליך.")
    add_heading(doc, "Redis Service", 2)
    add_para(doc, "ה-Service בשם redis-master בוחר Pods באמצעות labels. במצב התקין ה-selector הוא app redis ו-role master. כאשר ה-selector אינו תואם ל-labels של Redis, Kubernetes יוצר Service ללא endpoints. שם ה-DNS עדיין קיים, אבל אין backend שמקבל traffic.")
    add_heading(doc, "תרחיש ה-incident", 2)
    add_para(doc, "בתרחיש שנבדק, ה-selector השתנה ל-role unavailable וה-liveness probe הופנה אל readyz. השילוב יצר שני אפקטים: Redis נעלמה מאחורי ה-Service, וה-Pod החדש התחיל restart loop משום שה-liveness שלו הייתה תלויה ב-Redis. החקירה של KAOP זיהתה את שני השינויים.")

    add_heading(doc, "Kubernetes והרשאות")
    add_heading(doc, "משאבי היישום", 2)
    add_bullets(doc, [
        "Deployment guestbook עם שני replicas, readiness ו-liveness probes ומגבלות משאבים.",
        "Deployment redis-master עם Pod יחיד ו-labels שתואמים ל-Service התקין.",
        "Services עבור Guestbook ועבור Redis.",
        "namespace בשם kaop-demo שמפריד את הדמו משאר ה-cluster.",
    ])
    add_heading(doc, "ServiceAccount ו-RBAC", 2)
    add_para(doc, "ServiceAccount בשם kaop-investigator מחובר ל-Role בתוך namespace kaop-demo. ה-Role מאפשר get, list ו-watch בלבד עבור Pods, logs, events, Services, Endpoints, ConfigMaps, Deployments ו-ReplicaSets.")
    add_table(doc, ["פעולה", "תוצאה", "סיבה"], [
        ["קריאת Pods ולוגים", "Allowed", "נדרשת לאיסוף ראיות"],
        ["קריאת Services ו-Endpoints", "Allowed", "נדרשת לזיהוי selector mismatch"],
        ["קריאת Secrets", "Denied", "ה-resource אינו מופיע ב-Role"],
        ["patch או update", "Denied", "אין verbs של write"],
        ["גישה לכל ה-namespaces", "Denied", "Role ו-RoleBinding מוגבלים ל-kaop-demo"],
        ["pods exec", "Denied", "pods exec אינו כלול בהרשאות"],
    ], [5.0, 3.2, 8.0])
    add_para(doc, "המילה Denied משמשת כאן לתיאור גבול אבטחה שנאכף. היא אינה מציינת תקלה. דווקא התוצאה no בבדיקות ההרשאה מוכיחה שה-agent אינו יכול לחרוג מתפקיד החקירה.")

    add_heading(doc, "Grafana Cloud וה-alert")
    add_para(doc, "Grafana Alloy ו-kube-state-metrics מספקים את נתוני ה-cluster. כלל ה-alert עוקב אחר kube_pod_container_status_restarts_total עבור Pods של Guestbook ב-namespace kaop-demo. כאשר הערך גדול מאפס, הכלל עובר ל-Pending ולאחר חלון ההמתנה ל-Firing.")
    add_table(doc, ["מצב", "משמעות", "מה רואים בדמו"], [
        ["Normal", "התנאי אינו מתקיים", "אין restarts פעילים בבסיס התקין"],
        ["Pending", "התנאי מתקיים אך חלון ההמתנה טרם הסתיים", "restart זוהה והמערכת ממתינה דקה"],
        ["Firing", "התנאי מתקיים לאחר חלון ההמתנה", "Grafana שולחת webhook ל-KAOP"],
        ["Normal לאחר recovery", "התנאי חזר לערך תקין", "ה-alert נחשב resolved"],
    ], [3.3, 6.2, 6.7])
    add_para(doc, "ה-Contact Point שולח payload שמזהה deployment בשם guestbook, cluster בשם kaop-demo, namespace בשם kaop-demo ותקציב חקירה של 600 שניות. ה-token נשמר בממשק של Grafana ואינו נמצא ב-Git.")

    add_heading(doc, "KAOP ו-workflow החקירה")
    add_para(doc, "KAOP מקבלת את webhook ומפעילה את Kubernetes RCA agent. ה-agent משתמש בהרשאות של kaop-investigator כדי לבדוק Deployments, ReplicaSets, Pods, logs, events, Services ו-Endpoints. הוא אינו מבצע remediation.")
    add_para(doc, "ה-run המאומת run_e15c91fc04701d47c908ec0f הסתיים לאחר 2 דקות ו-48 שניות עם 15 tool calls. הוא מצא ש-Service redis-master בחר role unavailable בעוד Pod של Redis נשא role master. הוא גם השווה את ה-ReplicaSet החדש לישן וזיהה שה-liveness השתנתה מ-livez ל-readyz. ה-confidence היה high.")
    add_heading(doc, "Output contract", 2)
    add_table(doc, ["שדה", "תוכן"], [
        ["root cause", "הגורם הישיר והקשר בין שינוי התצורה להתנהגות המערכת"],
        ["evidence", "Services, Endpoints, Pod labels, logs, probes, events ו-restarts"],
        ["confidence", "low, medium או high לפי חוזק הראיות"],
        ["affected resources", "המשאבים שנפגעו בפועל"],
        ["suggested fix", "החזרת selector ל-role master והחזרת liveness ל-livez"],
    ], [4.2, 12.0])

    add_heading(doc, "GitHub ו-Code Reviewer")
    add_para(doc, "ה-repository הציבורי nirittere/kaop-guestbook-demo מכיל את האפליקציה, manifests, tests וחומרי ההגשה. ענף fix/redis-service-selector מכיל תיקון אמיתי: selector מתאים ל-Redis Pods, liveness חוזרת ל-livez, readiness נשארת ב-readyz וה-tests מאמתים את שני התנאים.")
    add_para(doc, "PR מספר 1 נשאר פתוח לצורך הדמו. AgentOps כתב review מסוג Looks good ובדק את שלושת הקבצים ששונו. ה-reviewer יכול לקרוא repository content, לקרוא PR ולכתוב review. הוא אינו מקבל כלי push, merge או שינוי קוד.")

    add_heading(doc, "חומרי ההגשה")
    add_table(doc, ["קובץ", "מטרה"], [
        ["KAOP-by-Nirit-Terehovsky.pptx", "מצגת עריכה עם 12 שקפים ו-speaker notes"],
        ["KAOP-by-Nirit-Terehovsky.pdf", "גרסה יציבה לצפייה ולגיבוי"],
        ["pitch-script.md", "תסריט דיבור קצר באנגלית"],
        ["demo-runbook.md", "סדר הפעולות לדמו החי"],
        ["demo-recovery.md", "תכנית התאוששות במקרה של תקלה"],
    ], [7.2, 9.0])

    add_heading(doc, "זרימת המידע והאבטחה")
    add_numbered(doc, [
        "ה-cluster מריץ את Guestbook ואת Redis ומייצר metrics ואירועים אמיתיים.",
        "Grafana Cloud מעריכה את כלל ה-alert על בסיס metrics של restarts.",
        "כאשר הכלל מגיע ל-Firing, Grafana שולחת webhook ל-KAOP.",
        "KAOP מפעילה RCA agent עם read only RBAC ב-namespace kaop-demo.",
        "ה-agent מחזיר finding מובנה ושומר run record עם הראיות וה-tool calls.",
        "GitHub מציג את תיקון התצורה ו-AgentOps משאיר review comment בלבד.",
        "המפעילה מריצה make recover ומוודאת ש-Grafana חוזרת ל-Normal.",
    ])
    add_para(doc, "אין tokens, credentials או Secret values ב-repository. Grafana מחזיקה את webhook credential בממשק שלה. Kubernetes מגביל את החקירה באמצעות RBAC. GitHub מגביל את ה-reviewer באמצעות הרשאות האפליקציה והכלים שהוקצו לו.")

    add_heading(doc, "קישורי ראיות")
    add_link_lines(doc, [
        ("Repository", "https://github.com/nirittere/kaop-guestbook-demo"),
        ("PR", "https://github.com/nirittere/kaop-guestbook-demo/pull/1"),
        ("AgentOps comment", "https://github.com/nirittere/kaop-guestbook-demo/pull/1#issuecomment-6045140326"),
        ("KAOP run", "https://kaop.komodor.com/a/hire-task-9/runs/run_e15c91fc04701d47c908ec0f"),
        ("Grafana alert", "https://contentcheetah182.grafana.net/alerting/grafana/fg0jdq07jv9c0e/view"),
    ])

    return save(doc, "02-KAOP-system-and-architecture-guide-he.docx")


def build_demo_guide():
    doc = configure_document(
        "מדריך הדמו של KAOP צעד אחר צעד",
        "פעולות תוצאות צפויות ונקודות הסבר במהלך ההדגמה",
    )
    add_contents(doc, [
        "הכנה מוקדמת",
        "שלב 1 הצגת KAOP וההרשאות",
        "שלב 2 אימות מצב תקין",
        "שלב 3 יצירת incident",
        "שלב 4 הצגת הראיות ב-Kubernetes",
        "שלב 5 Grafana ו-KAOP",
        "שלב 6 GitHub reviewer",
        "שלב 7 התאוששות",
        "תכנית גיבוי",
    ])

    add_heading(doc, "הכנה מוקדמת")
    add_para(doc, "יש לבצע את ההכנה כעשרים דקות לפני תחילת הראיון. המטרה היא להתחיל ממצב תקין, לוודא שכל הקישורים פתוחים ולהימנע מחשיפת credentials בזמן שיתוף המסך.")
    add_heading(doc, "הפעלת הסביבה", 2)
    add_code(doc, "make bootstrap\nmake healthy\nmake test")
    add_para(doc, "תוצאה צפויה: Colima פעילה, cluster בשם kaop-demo קיים, deployments של Guestbook ו-Redis מוכנים וכל הבדיקות עוברות ללא failures.")
    add_heading(doc, "פתיחת מסכים מראש", 2)
    add_bullets(doc, [
        "KAOP Kubernetes RCA agent וה-run history.",
        "Grafana alert בשם KAOP demo guestbook restart detected.",
        "GitHub PR מספר 1 ו-AgentOps review comment.",
        "Terminal פתוח ב-root של repository.",
    ])

    steps = [
        ("שלב 1 הצגת KAOP וההרשאות", "0:00 עד 2:00", [
            ("פעולה", "פתחי את Agent inventory ואת Kubernetes RCA agent ב-KAOP."),
            ("מה להראות", "ה-agent מחובר ל-cluster, ה-workflow מקבל trigger מ-Grafana וקיים run history."),
            ("מה לומר", "החקירה מוגבלת ל-namespace kaop-demo ולקריאה בלבד. Secrets, writes ו-pods exec חסומים."),
            ("תוצאה צפויה", "הקהל מבין את גבול ההרשאות לפני הצגת תוצאת החקירה."),
        ]),
        ("שלב 2 אימות מצב תקין", "2:00 עד 3:00", [
            ("פעולה", "הריצי make verify ובדקי את Redis endpoint."),
            ("פקודה", "make verify"),
            ("תוצאה צפויה", "guestbook מציג 2/2 Ready, redis-master מציג 1/1 Ready, ל-Service redis-master יש endpoint ו-restarts הם אפס."),
            ("מה לומר", "זהו baseline תקין. מכאן כל שינוי שיוצג נובע מה-incident שנפעיל."),
        ]),
        ("שלב 3 יצירת incident", "3:00 עד 4:00", [
            ("פעולה", "הריצי את פקודת ה-incident מתוך main או מסביבת הדמו שמכילה את overlay התקלה."),
            ("פקודה", "make incident"),
            ("תוצאה צפויה", "Kubernetes מחיל selector שגוי על Redis Service ומפעיל rollout של Guestbook עם liveness שתלויה ב-readiness."),
            ("מה לומר", "התקלה אינה log מלאכותי. היא שינוי Kubernetes אמיתי שמסיר endpoints וגורם ל-probe failures."),
        ]),
        ("שלב 4 הצגת הראיות ב-Kubernetes", "4:00 עד 6:00", [
            ("פעולה", "המתיני מספר שניות והריצי make verify. לאחר מכן הציגי logs של Guestbook."),
            ("פקודה", "make verify\nkubectl -n kaop-demo logs deployment/guestbook --tail=30"),
            ("תוצאה צפויה", "redis-master מציג ENDPOINTS none. ה-Pod החדש אינו Ready. מונה ה-restarts עולה. Events כוללים Unhealthy, Killing ו-BackOff. הלוגים כוללים connection refused או timeout ל-Redis."),
            ("מה לומר", "יש כאן שרשרת ראיות: selector ללא התאמה, Service ללא endpoint, כשל חיבור, probes שנכשלים ו-restarts."),
        ]),
        ("שלב 5 Grafana ו-KAOP", "6:00 עד 8:30", [
            ("פעולה", "עברי ל-Grafana והציגי את מצב ה-alert. לאחר Firing פתחי את run החדש ב-KAOP."),
            ("תוצאה צפויה", "Grafana עוברת מ-Normal ל-Pending ול-Firing. webhook מפעיל את Kubernetes RCA workflow."),
            ("מה להראות ב-KAOP", "root cause, evidence, confidence, affected resources ו-suggested fix. הציגי שה-agent ראה selector role unavailable מול Pod label role master."),
            ("ראיית גיבוי", "אם run חדש איטי, פתחי את run_e15c91fc04701d47c908ec0f. הוא כולל high confidence, שבעה restarts ו-15 tool calls."),
        ]),
        ("שלב 6 GitHub reviewer", "8:30 עד 10:30", [
            ("פעולה", "פתחי את PR מספר 1 בענף fix/redis-service-selector."),
            ("מה להראות", "ה-selector חוזר ל-role master. liveness חוזרת ל-livez. readiness נשארת ב-readyz. tests חדשים מאמתים את ההתנהגות."),
            ("תוצאה צפויה", "AgentOps comment מציג Looks good ומסביר שהשינויים עקביים עם הבדיקות."),
            ("מה לומר", "ה-reviewer יכול להגיב, אך אינו יכול לשנות קוד, לבצע push או merge. ה-PR נשאר פתוח כראיה חיה."),
        ]),
        ("שלב 7 התאוששות", "10:30 עד 12:00", [
            ("פעולה", "חזרי ל-terminal והריצי recovery ולאחריו verify."),
            ("פקודה", "make recover\nmake verify"),
            ("תוצאה צפויה", "Redis endpoint חוזר, Guestbook מציג 2/2 Ready, אין restart loop חדש ו-Grafana חוזרת ל-Normal."),
            ("מה לומר", "הדמו מסתיים במערכת תקינה וב-run record שניתן לבדיקה. לא בוצעה remediation אוטונומית."),
        ]),
    ]
    for title, timing, fields in steps:
        add_heading(doc, title)
        add_para(doc, f"זמן מומלץ: {timing}")
        for label, value in fields:
            if label == "פקודה":
                add_heading(doc, label, 2)
                add_code(doc, value)
            else:
                add_para(doc, f"{label}: {value}", f"{label}:")

    add_heading(doc, "בדיקות הרשאה להצגה לפי הצורך")
    add_code(doc, "kubectl auth can-i get pods --as=system:serviceaccount:kaop-demo:kaop-investigator -n kaop-demo\nkubectl auth can-i get secrets --as=system:serviceaccount:kaop-demo:kaop-investigator -n kaop-demo\nkubectl auth can-i patch deployments --as=system:serviceaccount:kaop-demo:kaop-investigator -n kaop-demo")
    add_para(doc, "תוצאה צפויה לפי הסדר: yes, no, no. המשמעות היא שה-agent יכול לאסוף ראיות אך אינו יכול לקרוא Secrets או לשנות Deployment.")

    add_heading(doc, "תכנית גיבוי")
    add_table(doc, ["בעיה", "פעולה", "ראיה חלופית"], [
        ["Grafana אינה מגיעה ל-Firing בזמן", "הציגי את כלל ה-alert ואת ה-evaluation האחרון", "ה-run המאומת ב-KAOP והראיות החיות ב-Kubernetes"],
        ["KAOP אינה מגיעה ל-cluster", "הציגי את run history ואת בדיקות RBAC", "make verify, logs ו-events"],
        ["AgentOps אינו מגיב מחדש", "אל תפתחי PR נוסף", "ה-comment הקיים ב-PR מספר 1"],
        ["ה-cluster אינו זמין", "בדקי colima status ו-k3d cluster list", "המצגת, ה-run השמור וקישורי הראיות"],
        ["נחשף מסך credentials", "עצרי שיתוף, סגרי את המסך וחזרי לטאב שהוכן מראש", "אין להציג token או webhook secret"],
    ], [4.4, 5.8, 6.0])

    add_heading(doc, "איפוס סופי לאחר הדמו")
    add_code(doc, "make recover\nkubectl -n kaop-demo get pods\nkubectl -n kaop-demo get endpoints redis-master")
    add_para(doc, "התוצאה הסופית הנדרשת היא Guestbook ו-Redis במצב Running ו-Ready, Redis endpoint קיים וה-alert במצב Normal.")

    add_heading(doc, "קישורים לפתיחה מראש")
    add_link_lines(doc, [
        ("KAOP agent", "https://kaop.komodor.com/a/hire-task-9/agents/agt_361674dea38c9ac19ee24b9a"),
        ("KAOP verified run", "https://kaop.komodor.com/a/hire-task-9/runs/run_e15c91fc04701d47c908ec0f"),
        ("Grafana alert", "https://contentcheetah182.grafana.net/alerting/grafana/fg0jdq07jv9c0e/view"),
        ("GitHub PR", "https://github.com/nirittere/kaop-guestbook-demo/pull/1"),
        ("AgentOps review", "https://github.com/nirittere/kaop-guestbook-demo/pull/1#issuecomment-6045140326"),
    ])

    return save(doc, "03-KAOP-live-demo-step-by-step-he.docx")


if __name__ == "__main__":
    outputs = [build_talk_guide(), build_system_guide(), build_demo_guide()]
    for path in outputs:
        print(path)
